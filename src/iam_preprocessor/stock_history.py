"""Auditable panels of observed surviving assets; no inferred disposal dates."""

from collections import defaultdict
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import yaml
from openpyxl import load_workbook

from .io import sha256_file
from .stock_observations import US_STATES, read_eia_generators


def registry():
    return yaml.safe_load((Path(__file__).parent / "data/remind_stock_assets.yaml").read_text())[
        "assets"
    ]


def number(value):
    if isinstance(value, (int, float)) and np.isfinite(value):
        return float(value)
    return None


def identifier(value):
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def selection_reason(row, mapping, *, statuses=("OP",), allow_chp=False, geography="USA"):
    if row["technology"] not in mapping["source_technologies"]:
        return "different_technology"
    if row["state"] not in US_STATES:
        return "outside_USA"
    if geography == "WECC" and row["nerc"] != "WECC":
        return "outside_WECC"
    if row["status"] not in statuses:
        return "non_operating"
    if not allow_chp and row["chp"] != "N":
        return "CHP_or_unknown"
    if mapping.get("required_fuel") and row["fuel"] != mapping["required_fuel"]:
        return "different_fuel"
    return "included"


def read_snapshot(path, year):
    """Preserve all status sheets, source identities and date/capacity metadata."""
    digest = sha256_file(path)
    with ZipFile(path) as archive:
        workbook = load_workbook(
            BytesIO(archive.read(f"2___Plant_Y{year}.xlsx")), read_only=True, data_only=True
        )
        rows = workbook["Plant"].iter_rows(values_only=True)
        next(rows)
        header = list(next(rows))
        regions = {}
        for values in rows:
            row = dict(zip(header, values, strict=True))
            code = row.get("Plant Code")
            if isinstance(code, (int, float)):
                if int(code) in regions:
                    raise ValueError("Duplicate EIA plant")
                regions[int(code)] = row.get("NERC Region")
        workbook.close()
    records, seen = [], set()
    for r in read_eia_generators(path, year, all_sheets=True):
        key = (int(r["Plant Code"]), identifier(r["Generator ID"]))
        if key in seen:
            raise ValueError(f"Duplicate generator across EIA sheets: {year}/{key}")
        seen.add(key)
        capacity = number(r.get("Nameplate Capacity (MW)"))
        if (capacity is None and r.get("Status") == "OP") or (
            capacity is not None and capacity < 0
        ):
            raise ValueError("Missing or negative generator capacity")
        cohort = number(r.get("Operating Year"))
        if cohort is not None and (cohort != int(cohort) or not 1850 <= cohort <= year):
            raise ValueError("Invalid operating year")
        records.append(
            dict(
                plant_id=key[0],
                generator_id=key[1],
                raw_generator_id=str(r["Generator ID"]),
                unit_id=identifier(r.get("Unit Code")),
                technology=r.get("Technology"),
                prime_mover=r.get("Prime Mover"),
                fuel=r.get("Energy Source 1"),
                chp=r.get("Associated with Combined Heat and Power System"),
                status=r.get("Status"),
                state=r.get("State"),
                nerc=regions.get(key[0]),
                capacity=capacity,
                summer_capacity=number(r.get("Summer Capacity (MW)")),
                winter_capacity=number(r.get("Winter Capacity (MW)")),
                cohort=int(cohort) if cohort is not None else None,
                operating_month=number(r.get("Operating Month")),
                retirement_year=number(r.get("Retirement Year")),
                retirement_month=number(r.get("Retirement Month")),
                uprate_derate=r.get("Uprate or Derate Completed During Year"),
                capacity_basis="nameplate_MW",
                cohort_date_kind="operating_year",
                observation_year=year,
                source_sha256=digest,
                source_member=r["source_member"],
                source_sheet=r["source_sheet"],
                source_row=r["source_row"],
            )
        )
    return records


def select_snapshot(records, mapping, *, statuses=("OP",), allow_chp=False, geography="USA"):
    candidates = [r for r in records if r["technology"] in mapping["source_technologies"]]
    ledger, chosen = [], []
    for row in candidates:
        reason = selection_reason(
            row, mapping, statuses=statuses, allow_chp=allow_chp, geography=geography
        )
        ledger.append(dict(row, selection_reason=reason))
        if reason == "included":
            chosen.append(row)
    return chosen, ledger


def summarize(records):
    cohorts = defaultdict(float)
    unknown = 0.0
    for r in records:
        if r["cohort"] is None:
            unknown += r["capacity"]
        else:
            cohorts[r["cohort"]] += r["capacity"]
    return dict(
        cohorts=dict(cohorts),
        unknown=unknown,
        total=sum(cohorts.values()) + unknown,
        generator_records=len(records),
    )


def transition_ledger(opening, closing, mapping, asset):
    """Disjoint contributions reconcile selected capacity without inventing retirements."""
    a = {(r["plant_id"], r["generator_id"]): r for r in opening}
    b = {(r["plant_id"], r["generator_id"]): r for r in closing}
    if len(a) != len(opening) or len(b) != len(closing):
        raise ValueError("Ambiguous panel identity")
    rows = []
    for key in sorted(a.keys() | b.keys()):
        old, new = a.get(key), b.get(key)
        was = old is not None and selection_reason(old, mapping) == "included"
        now = new is not None and selection_reason(new, mapping) == "included"
        if not was and not now:
            continue
        q0, q1 = (old["capacity"] if was else 0), (new["capacity"] if now else 0)
        if was and now:
            if old["cohort"] != new["cohort"]:
                event = "cohort_date_revision"
            elif q0 != q1:
                event = "capacity_change"
            else:
                event = "continuing"
        elif was:
            if new is None:
                event = "unexplained_disappearance"
            elif new["status"] == "RE":
                event = "reported_retirement"
            elif new["technology"] != old["technology"] or new["fuel"] != old["fuel"]:
                event = "category_change_out"
            else:
                event = "status_or_boundary_exit"
        else:
            if new["cohort"] is not None and new["cohort"] > opening[0]["observation_year"]:
                event = "dated_commissioning_entry"
            elif old is None:
                event = "newly_covered_old_asset"
            elif old["technology"] != new["technology"] or old["fuel"] != new["fuel"]:
                event = "category_change_in"
            else:
                event = "return_or_boundary_entry"
        rows.append(
            dict(
                asset=asset,
                start_year=opening[0]["observation_year"],
                end_year=closing[0]["observation_year"],
                plant_id=key[0],
                generator_id=key[1],
                opening_capacity=q0,
                closing_capacity=q1,
                capacity_delta=q1 - q0,
                event=event,
                opening_cohort=old["cohort"] if old else None,
                closing_cohort=new["cohort"] if new else None,
                opening_status=old["status"] if old else None,
                closing_status=new["status"] if new else None,
                opening_source_row=old["source_row"] if old else None,
                closing_source_row=new["source_row"] if new else None,
            )
        )
    actual = sum(r["capacity"] for r in b.values() if selection_reason(r, mapping) == "included")
    actual -= sum(r["capacity"] for r in a.values() if selection_reason(r, mapping) == "included")
    if not np.isclose(sum(r["capacity_delta"] for r in rows), actual, rtol=1e-10, atol=1e-8):
        raise ValueError("Transition ledger does not close")
    return rows
