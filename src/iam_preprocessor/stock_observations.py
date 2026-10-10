"""Auditable initial-stock curation for the bounded REMIND sector example.

Quantities are observed surviving stock, never historical inflows. Unknown dates
are retained in the audit and assigned proportionally only in the model input.
"""

from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import numpy as np
from openpyxl import load_workbook

NS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
}
US_STATES = set(
    "AL AK AZ AR CA CO CT DE DC FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN MS MO MT "
    "NE NV NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA WV WI WY".split()
)
POWER_TECHNOLOGIES = {
    "wind-onshore": {"Onshore Wind Turbine"},
    "wind-offshore": {"Offshore Wind Turbine"},
    "nuclear": {"Nuclear"},
    "gas-turbine": {"Natural Gas Fired Combustion Turbine"},
    "geothermal": {"Geothermal"},
    "csp": {"Solar Thermal without Energy Storage", "Solar Thermal with Energy Storage"},
}


def ods_rows(path, sheet):
    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("content.xml"))
    table = next(
        t for t in root.findall(".//table:table", NS) if t.get(f"{{{NS['table']}}}name") == sheet
    )
    for row in table.findall("table:table-row", NS):
        cells = []
        elements = row.findall("table:table-cell", NS)
        while (
            elements
            and not "".join(elements[-1].itertext())
            and elements[-1].get(f"{{{NS['office']}}}value") is None
        ):
            elements.pop()
        for cell in elements:
            repeat = int(cell.get(f"{{{NS['table']}}}number-columns-repeated", "1"))
            if repeat > 10000:
                raise ValueError("Unexpected repeated ODS column count")
            cells.extend(
                [("".join(cell.itertext()), cell.get(f"{{{NS['office']}}}value"))] * repeat
            )
        for _ in range(int(row.get(f"{{{NS['table']}}}number-rows-repeated", "1"))):
            yield cells


def vehicle_count(cell):
    display, raw = cell
    if raw is None:
        raise ValueError(f"Missing numeric vehicle count: {display!r}")
    value = Decimal(raw) * 1000
    count = int(value.to_integral_value())
    if count < 0 or abs(value - count) > Decimal("0.000001"):
        raise ValueError(f"Unexpected non-integer vehicle count: {raw!r}")
    return count


def curate_liquid_cars(path, year=2022):
    """Read disjoint petrol/diesel first-use cohorts; do not infer an IAM fuel split."""
    rows = list(ods_rows(path, "VEH1111"))
    start = next(i for i, r in enumerate(rows) if r and r[0][0].startswith("Geography"))
    header = [c[0] for c in rows[start]]
    columns = [(i, int(v)) for i, v in enumerate(header) if v.isdigit()]
    unknown = next(i for i, v in enumerate(header) if v.startswith("Unknown"))
    total = header.index("Total")
    fuels = {}
    for index, row in enumerate(rows[start + 1 :], start=start + 2):
        if len(row) < len(header):
            continue
        values = [c[0] for c in row]
        if (
            values[0] != "United Kingdom"
            or values[1] != f"{year} Q4 (end December)"
            or values[3] != "Cars"
            or values[4] not in {"PETROL", "DIESEL"}
        ):
            continue
        if values[2] != "Thousands" or values[4] in fuels:
            raise ValueError("Unexpected unit or duplicate selected fuel row")
        cohorts = {}
        for column, cohort in columns:
            if cohort > year:
                if row[column][0] != "[z]":
                    raise ValueError("Future observed vehicle cohort")
                continue
            count = vehicle_count(row[column])
            if count:
                cohorts[cohort] = count
        missing, stock = vehicle_count(row[unknown]), vehicle_count(row[total])
        if stock != sum(cohorts.values()) + missing:
            raise ValueError("Vehicle first-use quantities do not reconcile")
        fuels[values[4]] = dict(cohorts=cohorts, unknown=missing, total=stock, source_row=index)
    if set(fuels) != {"PETROL", "DIESEL"}:
        raise ValueError("Both petrol and diesel observations are required")
    combined = defaultdict(float)
    for row in fuels.values():
        for cohort, value in row["cohorts"].items():
            combined[cohort] += value
    return {
        "combined": dict(
            cohorts=dict(combined),
            unknown=sum(r["unknown"] for r in fuels.values()),
            total=sum(r["total"] for r in fuels.values()),
        ),
        "petrol": fuels["PETROL"],
        "diesel": fuels["DIESEL"],
    }


def read_eia_generators(path):
    with ZipFile(path) as archive:
        workbook = load_workbook(
            BytesIO(archive.read("3_1_Generator_Y2022.xlsx")), read_only=True, data_only=True
        )
    try:
        rows = workbook["Operable"].iter_rows(values_only=True)
        next(rows)
        header = [str(v).strip() if v is not None else "" for v in next(rows)]
        return [dict(zip(header, row, strict=True)) for row in rows if row[0] is not None]
    finally:
        workbook.close()


def curate_power(records, asset, year=2022):
    """Select operating US generators, exclude CHP, retain an exclusion mass ledger."""
    selected = [r for r in records if r["Technology"] in POWER_TECHNOLOGIES[asset]]
    seen, cohorts, excluded = set(), defaultdict(float), {}
    total, included, unknown = 0.0, 0, 0.0
    for r in selected:
        key = (r["Plant Code"], str(r["Generator ID"]))
        if key in seen:
            raise ValueError(f"Duplicate EIA generator: {key}")
        seen.add(key)
        capacity = float(r["Nameplate Capacity (MW)"])
        if not np.isfinite(capacity) or capacity <= 0:
            raise ValueError("Invalid EIA nameplate capacity")
        total += capacity
        reason = None
        if r["State"] not in US_STATES:
            reason = "outside_50_states_and_DC"
        elif r["Status"] != "OP":
            reason = "not_operating"
        elif r["Associated with Combined Heat and Power System"] != "N":
            reason = "CHP_or_unknown_CHP_status"
        if reason:
            entry = excluded.setdefault(reason, {"records": 0, "capacity_mw": 0.0})
            entry["records"] += 1
            entry["capacity_mw"] += capacity
            continue
        included += 1
        date = r["Operating Year"]
        if date in (None, "", " "):
            unknown += capacity
            continue
        date = float(date)
        if not np.isfinite(date) or date != int(date) or not 1872 <= date <= year:
            raise ValueError("Invalid EIA operating year")
        cohorts[int(date)] += capacity
    retained = sum(cohorts.values()) + unknown
    if not cohorts or not np.isclose(
        retained + sum(e["capacity_mw"] for e in excluded.values()), total, rtol=1e-12
    ):
        raise ValueError("EIA capacity selection is empty or does not reconcile")
    return dict(
        cohorts=dict(cohorts),
        unknown=unknown,
        total=retained,
        included_generator_records=included,
        selected_generator_records=len(selected),
        selected_capacity_mw=total,
        excluded=excluded,
    )


def model_observation(record, reference_year=2022):
    """Allocate unknown dates proportionally to observed shares, explicitly and once."""
    cohorts = {int(c): float(v) for c, v in record["cohorts"].items()}
    known, unknown, total = sum(cohorts.values()), record["unknown"], record["total"]
    if (
        not cohorts
        or known <= 0
        or unknown < 0
        or not np.isfinite(total)
        or any(v < 0 or not np.isfinite(v) for v in cohorts.values())
        or not np.isclose(known + unknown, total, rtol=1e-12)
        or min(cohorts) < reference_year - 150
        or max(cohorts) > reference_year
    ):
        raise ValueError("Invalid observed stock or unaccounted unknown dates")
    return [dict(cohort=c, amount=v * total / known) for c, v in sorted(cohorts.items())]
