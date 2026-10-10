"""Scientific failure modes for evidence, holdouts and service weights."""

from math import gamma

import numpy as np
import pandas as pd
import pytest

from iam_preprocessor.exceptions import DataValidationError
from iam_preprocessor.stock_history import select_snapshot, transition_ledger
from iam_preprocessor.stock_service import audit_service, match_ids
from iam_preprocessor.stock_usage import age_factors, read_age_profile
from iam_preprocessor.stock_validation import (
    check_experiment_inputs,
    conditional_history,
    opening_survival,
    score_cohorts,
)


def record(generator="1", cohort=2000, capacity=10, **extra):
    result = dict(
        plant_id=1,
        generator_id=generator,
        cohort=cohort,
        capacity=capacity,
        status="OP",
        technology="Wind",
        fuel="WND",
        prime_mover="WT",
        chp="N",
        state="CA",
        nerc="WECC",
        observation_year=2022,
        source_row=1,
        uprate_derate="N",
    )
    result.update(extra)
    return result


MAPPING = dict(source_technologies=["Wind"], required_fuel=None)


def audit(rows, generation, plant=None, previous=None):
    previous = previous or [dict(r, observation_year=2021) for r in rows]
    plants = (
        {(1, "WT", "WND"): dict(generation=plant, chp=["N"], source_rows=[3])}
        if plant is not None
        else {}
    )
    return audit_service(rows, rows, previous, generation, plants, MAPPING, "wind")


def gen(value, **extra):
    return dict(generation=value, prime_mover="WT", chp="N", source_row=3, **extra)


def test_aliases_are_one_to_one_and_duplicates_never_consumed():
    assert match_ids([record("01")], {(1, "1"): gen(10)})[(1, "01")] == (
        (1, "1"),
        "numeric_padding_alias",
    )
    matches = match_ids([record("01"), record("1")], {(1, "1"): gen(10)})
    assert matches[(1, "01")][0] is None
    assert matches[(1, "1")][0] == (1, "1")
    assert match_ids([record()], {(1, "1"): dict(ambiguous=True)})[(1, "1")][0] is None


def test_partial_year_never_becomes_observed_zero_generation():
    rows, _, _, d = audit([record(cohort=2022)], {(1, "1"): gen(0)})
    assert rows[0]["raw_generation"] == 0
    assert rows[0]["assigned_generation"] is None
    assert not d["coverage_passed"]


@pytest.mark.parametrize(
    "value,reason", [(-1, "negative_net_generation"), (90000, "physical_bound")]
)
def test_bad_generation_is_preserved_and_excluded(value, reason):
    rows, _, _, d = audit([record()], {(1, "1"): gen(value)})
    assert rows[0]["raw_generation"] == value
    assert rows[0]["assigned_generation"] is None
    # Without any usable output, the final reason makes the missing calibration explicit.
    assert rows[0]["reason"] in (reason, "no_positive_calibration_mean")
    assert d["observed_capacity_share"] == 0


def test_plant_allocations_close_without_claiming_observed_age_coverage():
    rows, rec, _, d = audit([record("1", 2000), record("2", 2010)], {}, plant=10000)
    assert sum(r["assigned_generation"] for r in rows) == 10000
    assert rec[0]["unassigned_generation"] == 0
    assert all(r["evidence"] == "plant_allocated" for r in rows)
    assert not d["coverage_passed"]
    rows, _, _, d = audit([record("1", 2000), record("2", 2000)], {}, plant=10000)
    assert d["coverage_passed"]
    assert all(r["evidence"] == "single_cohort_observed" for r in rows)


def test_material_old_bin_can_fail_otherwise_high_coverage():
    rows = [record("1", 2015, 94), record("2", 1980, 6)]
    _, _, _, d = audit(rows, {(1, "1"): gen(9400)})
    assert d["observed_capacity_share"] == 0.94
    assert not d["coverage_passed"]


def test_conflicting_plant_total_is_not_clipped():
    rows, rec, _, d = audit([record()], {(1, "1"): gen(1000)}, plant=500)
    assert rows[0]["raw_generation"] == 1000
    assert rows[0]["assigned_generation"] is None
    assert rec[0]["issue"] == "negative_plant_residual"
    assert not d["coverage_passed"]


def test_panel_reconciliation_distinguishes_coverage_capacity_and_retirement():
    before = [
        record("1", capacity=10, observation_year=2020),
        record("2", capacity=20, observation_year=2020),
        record("3", capacity=3, observation_year=2020),
    ]
    after = [
        record("1", capacity=12),
        record("2", capacity=20, status="RE"),
        record("4", 1990, capacity=5),
    ]
    rows = transition_ledger(before, after, MAPPING, "wind")
    assert {r["event"] for r in rows} == {
        "capacity_change",
        "reported_retirement",
        "unexplained_disappearance",
        "newly_covered_old_asset",
    }
    assert sum(r["capacity_delta"] for r in rows) == 17 - 33
    assert len(select_snapshot(after, MAPPING)[0]) == 2


def test_conditional_survival_and_no_retirement_benchmark():
    actual = opening_survival({2000: 10}, 2020, 2022, 30)[2000]
    scale = 30 / gamma(1 + 1 / 3)
    assert actual == pytest.approx(10 * np.exp(-((22 / scale) ** 3 - (20 / scale) ** 3)))
    assert opening_survival({2000: 10}, 2020, 2022, 30, method="no-retirement")[2000] == 10


def test_known_additions_and_stock_conditioning_are_separate_experiments():
    a = conditional_history(
        {2010: 10}, 2020, {2020: 10, 2021: 999, 2022: 555}, 30, additions={2021: 2, 2022: 3}
    )
    b = conditional_history(
        {2010: 10}, 2020, {2020: 10, 2021: 5, 2022: 5}, 30, additions={2021: 2, 2022: 3}
    )
    assert a == b  # supplied later aggregate values cannot affect experiment B
    assert a[2021][2021] == 2
    c = conditional_history({2010: 10}, 2020, {2020: 10, 2021: 9, 2022: 8}, 30)
    assert sum(c[2022].values()) == pytest.approx(8)
    scores = score_cohorts(c[2022], {2010: 7, 2022: 1}, 2022, total_is_input=True)
    assert scores["stock_error"] is None
    with pytest.raises(ValueError, match="Held-out"):
        check_experiment_inputs("C", ["initial_cohorts", "lifetime_prior", "closing_cohorts"])


def test_metrics_use_annual_age_distance_and_reject_future_mass():
    s = score_cohorts({2020: 1}, {2010: 1}, 2022, total_is_input=False)
    assert s["age_w1"] == 10
    assert s["mean_age_bias"] == -10
    assert s["total_variation"] == 1
    with pytest.raises(ValueError, match="future"):
        score_cohorts({2023: 1}, {2020: 1}, 2022, total_is_input=False)


def test_age_profile_evolves_initial_and_new_cohorts_with_explicit_tail(tmp_path):
    path = tmp_path / "usage.csv"
    pd.DataFrame({"age": [0, 1, 2], "factor": [1, 0.7, 0.2]}).to_csv(path, index=False)
    np.testing.assert_allclose(
        age_factors(path, [2020, 2021, 2023], np.array([2019, 2021])),
        [[0.7, 0], [0.2, 1], [0.2, 0.2]],
    )
    pd.DataFrame({"age": [0, 2], "factor": [1, 0.2]}).to_csv(path, index=False)
    with pytest.raises(DataValidationError):
        read_age_profile(path)


@pytest.mark.parametrize("year", [2020, 2021, 2022])
def test_eia_year_reader_keeps_status_sheets_and_missing_capacity(tmp_path, year):
    from io import BytesIO
    from zipfile import ZipFile

    from openpyxl import Workbook

    from iam_preprocessor.stock_history import read_snapshot

    g = Workbook()
    g.remove(g.active)
    headers = ["Plant Code", "Generator ID", "Status", "Nameplate Capacity (MW)", "Operating Year"]
    for name, values in [
        ("Operable", [1, "01", "OP", 10, 2000]),
        ("Proposed", [1, "02", "CN", None, None]),
    ]:
        sheet = g.create_sheet(name)
        sheet.append(["Title"])
        sheet.append(headers)
        sheet.append(values)
    p = Workbook()
    sheet = p.active
    sheet.title = "Plant"
    sheet.append(["Title"])
    sheet.append(["Plant Code", "NERC Region"])
    sheet.append([1, "WECC"])
    path = tmp_path / f"{year}.zip"
    with ZipFile(path, "w") as archive:
        for name, workbook in [(f"3_1_Generator_Y{year}.xlsx", g), (f"2___Plant_Y{year}.xlsx", p)]:
            buffer = BytesIO()
            workbook.save(buffer)
            archive.writestr(name, buffer.getvalue())
    rows = read_snapshot(path, year)
    assert [r["generator_id"] for r in rows] == ["01", "02"]
    assert rows[1]["capacity"] is None and rows[1]["cohort"] is None
    assert rows[0]["nerc"] == "WECC"
    assert rows[0]["source_row"] == 3
    assert rows[0]["source_sha256"] == rows[1]["source_sha256"]


def test_generation_reader_preserves_conflicting_identity_and_missing_output(tmp_path):
    from io import BytesIO
    from zipfile import ZipFile

    from openpyxl import Workbook

    from iam_preprocessor.stock_service import read_generation

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Page 4 Generator Data"
    sheet.append(
        [
            "Plant Id",
            "Generator Id",
            "Reported Prime Mover",
            "Combined Heat And Power Plant",
            "Net Generation Year To Date",
        ]
    )
    sheet.append([1, "A", "WT", "N", 0])
    sheet.append([1, "A", "WT", "N", 5])
    sheet.append([1, "B", "WT", "N", "."])
    sheet = workbook.create_sheet("Page 1 Generation and Fuel Data")
    sheet.append(
        [
            "Plant Id",
            "Nuclear Unit Id",
            "Reported Prime Mover",
            "Reported Fuel Type Code",
            "Net Generation (Megawatthours)",
            "Combined Heat And Power Plant",
        ]
    )
    sheet.append([1, ".", "WT", "WND", 5, "N"])
    buffer = BytesIO()
    workbook.save(buffer)
    path = tmp_path / "generation.zip"
    with ZipFile(path, "w") as archive:
        archive.writestr("Schedules_2_3_4_5_2022.xlsx", buffer.getvalue())
    g, p = read_generation(path)
    assert g[1, "A"]["ambiguous"]
    assert [r["generation"] for r in g[1, "A"]["duplicate_rows"]] == [0, 5]
    assert g[1, "B"]["generation"] is None
    assert p[1, "WT", "WND"]["generation"] == 5
